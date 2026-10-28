"""Excel report generator for Finance Hub."""
import io
from datetime import datetime, date
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side, numbers
    )
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.series import DataPoint
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False


# ── Colour palette ────────────────────────────────────────────────────────────
NAVY   = '1e293b'
BLUE   = '3b82f6'
GREEN  = '22c55e'
RED    = 'ef4444'
AMBER  = 'f59e0b'
LIGHT  = 'f1f5f9'
WHITE  = 'FFFFFF'
GREY   = '64748b'


def _hdr(cell, text, bold=True, bg=NAVY, fg=WHITE, size=11):
    cell.value = text
    cell.font = Font(bold=bold, color=fg, size=size)
    cell.fill = PatternFill('solid', fgColor=bg)
    cell.alignment = Alignment(horizontal='center', vertical='center')


def _border():
    thin = Side(style='thin', color='cbd5e1')
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _money(ws, cell_ref, value):
    ws[cell_ref].value = value
    ws[cell_ref].number_format = u'₹#,##0.00'


def generate_monthly_report(transactions: list[dict], fy_start: str, fy_end: str) -> bytes:
    """Generate monthly summary Excel report. Returns bytes."""
    if not OPENPYXL_OK:
        return b''

    wb = openpyxl.Workbook()

    # ── Sheet 1: Monthly Summary ──────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Monthly Summary'

    # Title
    ws.merge_cells('A1:H1')
    ws['A1'].value = f'Personal Finance Report  ·  FY {fy_start[:4]}-{fy_end[2:4]}'
    ws['A1'].font = Font(bold=True, size=14, color=WHITE)
    ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    # Subtitle
    ws['A2'] = f'Generated on {datetime.now().strftime("%d %b %Y %H:%M")}'
    ws['A2'].font = Font(italic=True, color=GREY, size=9)
    ws['A3'] = ''

    # Headers
    headers = ['Month', 'Debit (Spend)', 'Credit (Received)', 'Net', 'Transactions', 'Top Category']
    for col, h in enumerate(headers, 1):
        _hdr(ws.cell(4, col), h, bg=BLUE)

    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 20

    # Aggregate by month
    from collections import defaultdict
    monthly: dict = defaultdict(lambda: {'debit': 0, 'credit': 0, 'count': 0, 'cats': defaultdict(float)})
    for t in transactions:
        try:
            d = datetime.strptime(t['date'], '%Y-%m-%d')
            if not (fy_start <= t['date'] <= fy_end):
                continue
            key = d.strftime('%b %Y')
            bucket = monthly[key]
            if t['type'] == 'debit':
                bucket['debit'] += t['amount']
                bucket['cats'][t.get('category', 'Others')] += t['amount']
            else:
                bucket['credit'] += t['amount']
            bucket['count'] += 1
        except Exception:
            pass

    # Sort months chronologically
    def month_sort(k):
        try:
            return datetime.strptime(k, '%b %Y')
        except Exception:
            return datetime.min

    sorted_months = sorted(monthly.keys(), key=month_sort)

    row = 5
    total_debit = total_credit = total_count = 0
    debit_series = []
    credit_series = []
    month_labels = []

    for mon in sorted_months:
        b = monthly[mon]
        top_cat = max(b['cats'], key=b['cats'].get) if b['cats'] else '—'
        net = b['credit'] - b['debit']

        ws.cell(row, 1).value = mon
        ws.cell(row, 2).value = round(b['debit'], 2)
        ws.cell(row, 3).value = round(b['credit'], 2)
        ws.cell(row, 4).value = round(net, 2)
        ws.cell(row, 5).value = b['count']
        ws.cell(row, 6).value = top_cat

        ws.cell(row, 2).number_format = u'₹#,##0.00'
        ws.cell(row, 3).number_format = u'₹#,##0.00'
        ws.cell(row, 4).number_format = u'₹#,##0.00'

        # Colour net positive green, negative red
        if net >= 0:
            ws.cell(row, 4).font = Font(color='16a34a', bold=True)
        else:
            ws.cell(row, 4).font = Font(color=RED, bold=True)

        # Zebra shading
        if row % 2 == 0:
            for col in range(1, 7):
                ws.cell(row, col).fill = PatternFill('solid', fgColor=LIGHT)

        for col in range(1, 7):
            ws.cell(row, col).border = _border()

        debit_series.append(round(b['debit'], 2))
        credit_series.append(round(b['credit'], 2))
        month_labels.append(mon)
        total_debit += b['debit']
        total_credit += b['credit']
        total_count += b['count']
        row += 1

    # Totals row
    for col in range(1, 7):
        ws.cell(row, col).fill = PatternFill('solid', fgColor=NAVY)
        ws.cell(row, col).font = Font(bold=True, color=WHITE)
        ws.cell(row, col).border = _border()
    ws.cell(row, 1).value = 'TOTAL'
    ws.cell(row, 2).value = round(total_debit, 2)
    ws.cell(row, 3).value = round(total_credit, 2)
    ws.cell(row, 4).value = round(total_credit - total_debit, 2)
    ws.cell(row, 5).value = total_count
    ws.cell(row, 2).number_format = u'₹#,##0.00'
    ws.cell(row, 3).number_format = u'₹#,##0.00'
    ws.cell(row, 4).number_format = u'₹#,##0.00'

    # ── Sheet 2: Category Analysis ────────────────────────────────────────────
    ws2 = wb.create_sheet('Category Analysis')
    ws2.merge_cells('A1:E1')
    ws2['A1'].value = 'Category-wise Expense Analysis'
    ws2['A1'].font = Font(bold=True, size=13, color=WHITE)
    ws2['A1'].fill = PatternFill('solid', fgColor=NAVY)
    ws2['A1'].alignment = Alignment(horizontal='center')
    ws2.row_dimensions[1].height = 28

    cat_headers = ['Category', 'Total Spent', 'Transactions', '% of Total', 'Avg per Txn']
    for col, h in enumerate(cat_headers, 1):
        _hdr(ws2.cell(3, col), h, bg=BLUE)

    ws2.column_dimensions['A'].width = 22
    ws2.column_dimensions['B'].width = 16
    ws2.column_dimensions['C'].width = 14
    ws2.column_dimensions['D'].width = 12
    ws2.column_dimensions['E'].width = 16

    cat_totals: dict = defaultdict(lambda: {'amount': 0, 'count': 0})
    for t in transactions:
        if t['type'] == 'debit' and fy_start <= t['date'] <= fy_end:
            cat = t.get('category', 'Others')
            cat_totals[cat]['amount'] += t['amount']
            cat_totals[cat]['count'] += 1

    sorted_cats = sorted(cat_totals.items(), key=lambda x: x[1]['amount'], reverse=True)
    grand = sum(v['amount'] for _, v in sorted_cats) or 1

    row2 = 4
    for cat, v in sorted_cats:
        pct = v['amount'] / grand * 100
        avg = v['amount'] / v['count'] if v['count'] else 0

        ws2.cell(row2, 1).value = cat
        ws2.cell(row2, 2).value = round(v['amount'], 2)
        ws2.cell(row2, 2).number_format = u'₹#,##0.00'
        ws2.cell(row2, 3).value = v['count']
        ws2.cell(row2, 4).value = round(pct, 1)
        ws2.cell(row2, 4).number_format = '0.0"%"'
        ws2.cell(row2, 5).value = round(avg, 2)
        ws2.cell(row2, 5).number_format = u'₹#,##0.00'

        if row2 % 2 == 0:
            for col in range(1, 6):
                ws2.cell(row2, col).fill = PatternFill('solid', fgColor=LIGHT)
        for col in range(1, 6):
            ws2.cell(row2, col).border = _border()
        row2 += 1

    # ── Sheet 3: Account Summary ──────────────────────────────────────────────
    ws3 = wb.create_sheet('Account Summary')
    ws3.merge_cells('A1:F1')
    ws3['A1'].value = 'Account-wise Summary'
    ws3['A1'].font = Font(bold=True, size=13, color=WHITE)
    ws3['A1'].fill = PatternFill('solid', fgColor=NAVY)
    ws3['A1'].alignment = Alignment(horizontal='center')
    ws3.row_dimensions[1].height = 28

    acc_headers = ['Account', 'Total Debit', 'Total Credit', 'Net', 'Transactions', 'Last Date']
    for col, h in enumerate(acc_headers, 1):
        _hdr(ws3.cell(3, col), h, bg=BLUE)
    for col, w in zip(range(1, 7), [22, 16, 16, 14, 14, 14]):
        ws3.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

    acc_totals: dict = defaultdict(lambda: {'debit': 0, 'credit': 0, 'count': 0, 'last': ''})
    for t in transactions:
        if fy_start <= t['date'] <= fy_end:
            a = t.get('account', 'Unknown')
            b2 = acc_totals[a]
            if t['type'] == 'debit':
                b2['debit'] += t['amount']
            else:
                b2['credit'] += t['amount']
            b2['count'] += 1
            if t['date'] > b2['last']:
                b2['last'] = t['date']

    row3 = 4
    for acc, v in sorted(acc_totals.items(), key=lambda x: x[1]['debit'], reverse=True):
        ws3.cell(row3, 1).value = acc
        ws3.cell(row3, 2).value = round(v['debit'], 2)
        ws3.cell(row3, 2).number_format = u'₹#,##0.00'
        ws3.cell(row3, 3).value = round(v['credit'], 2)
        ws3.cell(row3, 3).number_format = u'₹#,##0.00'
        net = v['credit'] - v['debit']
        ws3.cell(row3, 4).value = round(net, 2)
        ws3.cell(row3, 4).number_format = u'₹#,##0.00'
        ws3.cell(row3, 5).value = v['count']
        ws3.cell(row3, 6).value = v['last']
        if row3 % 2 == 0:
            for col in range(1, 7):
                ws3.cell(row3, col).fill = PatternFill('solid', fgColor=LIGHT)
        for col in range(1, 7):
            ws3.cell(row3, col).border = _border()
        row3 += 1

    # ── Sheet 4: All Transactions ─────────────────────────────────────────────
    ws4 = wb.create_sheet('All Transactions')
    txn_headers = ['Date', 'Description', 'Amount', 'Type', 'Category', 'Account']
    for col, h in enumerate(txn_headers, 1):
        _hdr(ws4.cell(1, h := col), h, bg=BLUE)
        ws4.cell(1, col).value = txn_headers[col - 1]
    for col, w in zip(range(1, 7), [12, 40, 14, 8, 20, 22]):
        ws4.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w
    ws4.freeze_panes = 'A2'

    sorted_txns = sorted(
        [t for t in transactions if fy_start <= t['date'] <= fy_end],
        key=lambda x: x['date']
    )
    for row4, t in enumerate(sorted_txns, 2):
        ws4.cell(row4, 1).value = t['date']
        ws4.cell(row4, 2).value = t.get('description', '')
        ws4.cell(row4, 3).value = round(t['amount'], 2)
        ws4.cell(row4, 3).number_format = u'₹#,##0.00'
        ws4.cell(row4, 4).value = t.get('type', '')
        ws4.cell(row4, 5).value = t.get('category', '')
        ws4.cell(row4, 6).value = t.get('account', '')
        if t.get('type') == 'credit':
            ws4.cell(row4, 3).font = Font(color='16a34a')
        if row4 % 2 == 0:
            for col in range(1, 7):
                ws4.cell(row4, col).fill = PatternFill('solid', fgColor=LIGHT)

    # ── Output ────────────────────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
