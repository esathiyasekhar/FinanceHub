"""Export StatementData to a formatted Excel workbook."""
import re
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                                  numbers)
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False


# ── Colour palette ──────────────────────────────────────────────────────────
C_HEADER_BG  = '1E3A5F'   # dark navy
C_HEADER_FG  = 'FFFFFF'
C_TITLE_BG   = '2E86AB'   # teal
C_TITLE_FG   = 'FFFFFF'
C_DEBIT      = 'FFE5E5'   # light red fill
C_CREDIT     = 'E5F5E5'   # light green fill
C_DEBIT_FG   = 'C0392B'
C_CREDIT_FG  = '1A7A3F'
C_ALT_ROW    = 'F7F9FC'
C_SECTION_BG = 'EBF4FF'


def _hdr(bold=True, fg=C_HEADER_FG, bg=C_HEADER_BG, sz=10):
    f = Font(bold=bold, color=fg, size=sz, name='Calibri')
    fill = PatternFill('solid', fgColor=bg)
    align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    return f, fill, align


def _auto_width(ws, min_w=10, max_w=50):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or '')))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(min_w, min(max_len + 2, max_w))


def _thin_border():
    s = Side(style='thin', color='D0D0D0')
    return Border(left=s, right=s, top=s, bottom=s)


def export_to_excel(data, out_path: str) -> str:
    """
    Export StatementData to a multi-sheet Excel file.
    Returns the path written.
    """
    if not OPENPYXL_OK:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    _write_summary(wb, data)
    if getattr(data, 'accounts', None):
        _write_accounts(wb, data)
    _write_transactions(wb, data)
    if data.emis:
        _write_emis(wb, data)
    if any([data.reward_opening, data.reward_earned, data.reward_redeemed,
            data.reward_closing, data.reward_details]):
        _write_rewards(wb, data)

    wb.save(out_path)
    return out_path


# ── Sheet writers ───────────────────────────────────────────────────────────

def _write_summary(wb, data):
    ws = wb.create_sheet('Summary')
    ws.sheet_view.showGridLines = False

    font_hdr, fill_hdr, align_ctr = _hdr()
    title_font = Font(bold=True, size=14, color=C_TITLE_FG, name='Calibri')
    title_fill = PatternFill('solid', fgColor=C_TITLE_BG)
    lbl_font   = Font(bold=True, size=10, name='Calibri')
    val_font   = Font(size=10, name='Calibri')

    # Title row
    ws.merge_cells('A1:D1')
    ws['A1'] = f"Statement Summary — {data.card_name or data.bank}"
    ws['A1'].font = title_font
    ws['A1'].fill = title_fill
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    rows = [
        ('Bank', data.bank),
        ('Card / Account', data.card_name or data.account_number),
        ('Card Number', data.card_number or data.account_number),
        ('Statement Type', data.statement_type.replace('_', ' ').title()),
        ('Source File', data.source_file),
        ('', ''),
        ('Statement Date', data.statement_date),
        ('Payment Due Date', data.due_date),
        ('Statement Period', f"{data.from_date} to {data.to_date}" if data.from_date else ''),
        ('', ''),
        ('Total Amount Due', data.total_due),
        ('Minimum Amount Due', data.minimum_due),
        ('Credit Limit', data.credit_limit),
        ('Available Credit', data.available_credit),
        ('Opening Balance', data.opening_balance),
        ('Closing Balance', data.closing_balance),
        ('', ''),
        ('Total Transactions', len(data.transactions)),
        ('Total Spent (Debits)', sum(t.amount for t in data.transactions if t.txn_type == 'debit')),
        ('Total Received (Credits)', sum(t.amount for t in data.transactions if t.txn_type == 'credit')),
        ('', ''),
        ('Reward Points — Opening', data.reward_opening),
        ('Reward Points — Earned', data.reward_earned),
        ('Reward Points — Redeemed', data.reward_redeemed),
        ('Reward Points — Expired', data.reward_expired),
        ('Reward Points — Closing', data.reward_closing),
        ('Reward Points Expiry', data.reward_expiry_date),
    ]

    # Labels follow the kind of document (the values are the same fields)
    st = (data.statement_type or '').lower()
    relabel = {}
    if st == 'consolidated':
        relabel = {'Total Amount Due': 'Total Loans Outstanding', 'Closing Balance': 'Total Deposit Balance',
                   'Card Number': 'Customer ID'}
    elif st == 'emi_letter':
        relabel = {'Total Amount Due': 'Loan Outstanding', 'Statement Date': 'Letter Date',
                   'Card / Account': 'Document'}
    elif 'loan' in st:
        relabel = {'Total Amount Due': 'Loan Outstanding (as on header date)', 'Minimum Amount Due': 'EMI',
                   'Credit Limit': 'Sanctioned Amount', 'Card Number': 'CIF / Customer No.'}
    elif st in ('savings', 'current'):
        relabel = {'Card Number': 'Customer ID'}
    rows = [(relabel.get(l, l), v) for l, v in rows]

    for r_idx, (lbl, val) in enumerate(rows, start=2):
        ws.cell(r_idx, 1, lbl).font = lbl_font
        ws.cell(r_idx, 1).fill = PatternFill('solid', fgColor=C_SECTION_BG if lbl else 'FFFFFF')
        cell = ws.cell(r_idx, 2, val if val != 0.0 else ('' if isinstance(val, float) else val))
        cell.font = val_font
        if isinstance(val, float) and val > 0:
            cell.number_format = '₹#,##0.00'
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 40

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 40


def _write_accounts(wb, data):
    """One row per account (consolidated statements) with a balance reconciliation check."""
    ws = wb.create_sheet('Accounts')
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'
    headers = list(data.accounts[0].keys())
    font_hdr, fill_hdr, align_ctr = _hdr()
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font, cell.fill, cell.alignment, cell.border = font_hdr, fill_hdr, align_ctr, _thin_border()
    money = {'Opening', 'Credits', 'Debits', 'Closing'}
    for r, acc in enumerate(data.accounts, 2):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(r, c, acc.get(h, ''))
            cell.border = _thin_border()
            cell.font = Font(size=10, name='Calibri')
            if h in money:
                cell.number_format = '#,##0.00'
            if h == 'Check':
                ok = str(acc.get(h, '')).startswith('OK')
                cell.font = Font(size=10, bold=True, name='Calibri', color='15803D' if ok else 'B91C1C')
    _auto_width(ws)


def _write_transactions(wb, data):
    ws = wb.create_sheet('Transactions')
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'

    headers = ['#', 'Date', 'Description', 'Type', 'Amount (₹)',
               'Reward Pts', 'Reference', 'Card/Account', 'Category', 'Full narration']
    font_hdr, fill_hdr, align_ctr = _hdr()

    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = font_hdr
        cell.fill = fill_hdr
        cell.alignment = align_ctr
        cell.border = _thin_border()
    ws.row_dimensions[1].height = 22

    debit_fill  = PatternFill('solid', fgColor=C_DEBIT)
    credit_fill = PatternFill('solid', fgColor=C_CREDIT)

    for i, txn in enumerate(data.transactions, 1):
        r = i + 1
        row_fill = debit_fill if txn.txn_type == 'debit' else credit_fill
        amt_fg    = C_DEBIT_FG if txn.txn_type == 'debit' else C_CREDIT_FG

        vals = [i, txn.date, txn.description,
                'Debit' if txn.txn_type == 'debit' else 'Credit',
                txn.amount, txn.reward_pts or '', txn.reference,
                txn.cardholder or data.card_name, txn.category,
                getattr(txn, 'narration', '') or '']

        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.fill = row_fill if c != 5 else PatternFill('solid',
                        fgColor=C_DEBIT if txn.txn_type == 'debit' else C_CREDIT)
            cell.border = _thin_border()
            cell.alignment = Alignment(vertical='center',
                                        horizontal='right' if c in (1, 5, 6) else 'left',
                                        wrap_text=(c == 3))
            if c == 5:
                cell.number_format = '#,##0.00'
                cell.font = Font(bold=True, color=amt_fg, size=10, name='Calibri')
            else:
                cell.font = Font(size=10, name='Calibri')

    # Column widths
    widths = [4, 12, 42, 8, 14, 10, 20, 22, 16, 60]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    # Totals row
    if data.transactions:
        r = len(data.transactions) + 2
        total_debit  = sum(t.amount for t in data.transactions if t.txn_type == 'debit')
        total_credit = sum(t.amount for t in data.transactions if t.txn_type == 'credit')
        ws.cell(r, 3, 'TOTAL').font = Font(bold=True, size=10, name='Calibri')
        ws.cell(r, 3).fill = PatternFill('solid', fgColor=C_HEADER_BG)
        ws.cell(r, 3).font = Font(bold=True, color='FFFFFF', size=10)
        c5 = ws.cell(r, 5, f'↑{total_credit:,.2f}  ↓{total_debit:,.2f}')
        c5.font = Font(bold=True, size=10, name='Calibri')


def _write_emis(wb, data):
    ws = wb.create_sheet('EMI Details')
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'

    headers = ['#', 'Description', 'Principal (₹)', 'Interest (₹)',
               'EMI Amount (₹)', 'Tenure (Months)', 'Months Remaining',
               'Start Date', 'End Date', 'Loan / Ref No.', 'Lender',
               'Original Amount (₹)', 'Outstanding (₹)', 'Rate % p.a.']
    font_hdr, fill_hdr, align_ctr = _hdr()
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = font_hdr
        cell.fill = fill_hdr
        cell.alignment = align_ctr
        cell.border = _thin_border()

    for i, emi in enumerate(data.emis, 1):
        r = i + 1
        alt = PatternFill('solid', fgColor=C_ALT_ROW) if i % 2 == 0 else None
        vals = [i, emi.description, emi.principal, emi.interest,
                emi.emi_amount, emi.tenure_months, emi.months_remaining,
                emi.start_date, emi.end_date, getattr(emi, 'ref', ''), getattr(emi, 'lender', ''),
                getattr(emi, 'original_amount', 0) or '', getattr(emi, 'outstanding', 0) or '',
                getattr(emi, 'rate_pa', 0) or '']
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            if alt:
                cell.fill = alt
            cell.border = _thin_border()
            if c in (3, 4, 5, 12, 13):
                cell.number_format = '#,##0.00'

    _auto_width(ws)


def _write_rewards(wb, data):
    ws = wb.create_sheet('Reward Points')
    ws.sheet_view.showGridLines = False

    font_hdr, fill_hdr, align_ctr = _hdr()
    lbl_font = Font(bold=True, size=10, name='Calibri')
    val_font = Font(size=10, name='Calibri')

    # Summary
    ws['A1'] = 'Reward Points Summary'
    ws['A1'].font = Font(bold=True, size=12, color=C_TITLE_FG, name='Calibri')
    ws['A1'].fill = PatternFill('solid', fgColor=C_TITLE_BG)
    ws.merge_cells('A1:B1')

    summary = [
        ('Opening Balance', data.reward_opening),
        ('Points Earned This Period', data.reward_earned),
        ('Points Redeemed', data.reward_redeemed),
        ('Points Expired', data.reward_expired),
        ('Closing Balance', data.reward_closing),
        ('Points Expiry Date', data.reward_expiry_date),
    ]
    for r_idx, (lbl, val) in enumerate(summary, 2):
        ws.cell(r_idx, 1, lbl).font = lbl_font
        ws.cell(r_idx, 1).fill = PatternFill('solid', fgColor=C_SECTION_BG)
        ws.cell(r_idx, 2, val).font = val_font

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 20

    # Detailed reward points table if available
    if data.reward_details:
        start_row = len(summary) + 4
        headers = ['Date', 'Description', 'Points Earned', 'Points Redeemed',
                   'Points Expired', 'Balance']
        for c, h in enumerate(headers, 1):
            cell = ws.cell(start_row, c, h)
            cell.font = font_hdr
            cell.fill = fill_hdr
            cell.alignment = align_ctr

        for i, rp in enumerate(data.reward_details, 1):
            r = start_row + i
            vals = [rp.date, rp.description, rp.points_earned,
                    rp.points_redeemed, rp.points_expired, rp.balance]
            for c, v in enumerate(vals, 1):
                ws.cell(r, c, v)
