"""Finance Hub 1.2 — Excel and PDF reports for whatever you are viewing (all / a book / an account / a broker).

Every report is a list of sections (key-value summaries and tables) rendered either as an Excel workbook
(one sheet per section) or as a landscape A4 PDF with page numbers.
"""
import io
import re
from collections import defaultdict
from datetime import date, datetime

from flask import request, send_file, flash, redirect, url_for

import classifier as C
import scope as SC

MONEY = 'money'


# ── formatting ──────────────────────────────────────────────────────────────
def inr(v, sym='₹'):
    try:
        v = float(v)
    except (TypeError, ValueError):
        return ''
    neg = v < 0
    v = abs(v)
    whole, frac = f'{v:.2f}'.split('.')
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        head = re.sub(r'(\d)(?=(\d{2})+$)', r'\1,', head)
        whole = f'{head},{tail}'
    return f"{'-' if neg else ''}{sym}{whole}.{frac}"


def fy_default():
    t = date.today()
    y = t.year if t.month >= 4 else t.year - 1
    return f'{y}-04-01', f'{y + 1}-03-31'


# ── renderers ───────────────────────────────────────────────────────────────
def to_xlsx(rep):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    head = PatternFill('solid', fgColor='4F46E5')
    thin = Side(style='thin', color='D8DCEB')
    used = set()
    for sec in rep['sections']:
        name = re.sub(r'[\\/*?:\[\]]', '', sec['title'])[:28] or 'Sheet'
        base, i = name, 2
        while name in used:
            name = f'{base[:25]} {i}'; i += 1
        used.add(name)
        ws = wb.create_sheet(name)
        ws['A1'] = rep['title']
        ws['A1'].font = Font(bold=True, size=14, color='312E81')
        ws['A2'] = rep.get('subtitle', '')
        ws['A2'].font = Font(italic=True, color='6B7280')
        ws['A3'] = sec['title']
        ws['A3'].font = Font(bold=True, size=12)
        r0 = 5
        if sec['kind'] == 'kv':
            for i, (k, v) in enumerate(sec['rows']):
                ws.cell(r0 + i, 1, k).font = Font(bold=True)
                c = ws.cell(r0 + i, 2, v)
                if isinstance(v, (int, float)):
                    c.number_format = '#,##,##0.00'
            ws.column_dimensions['A'].width = 38
            ws.column_dimensions['B'].width = 24
            continue
        cols = sec['columns']
        money = set(sec.get('money', []))
        for j, col in enumerate(cols, 1):
            c = ws.cell(r0, j, col)
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = head
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for i, row in enumerate(sec['rows'], 1):
            for j, v in enumerate(row, 1):
                c = ws.cell(r0 + i, j, v)
                c.border = Border(bottom=thin)
                if cols[j - 1] in money and isinstance(v, (int, float)):
                    c.number_format = '#,##,##0.00;[Red]-#,##,##0.00'
        if sec.get('totals'):
            tr = r0 + len(sec['rows']) + 1
            ws.cell(tr, 1, 'Total').font = Font(bold=True)
            for j, col in enumerate(cols, 1):
                if col in sec['totals']:
                    c = ws.cell(tr, j, sum(float(r[j - 1] or 0) for r in sec['rows'] if isinstance(r[j - 1], (int, float))))
                    c.font = Font(bold=True)
                    c.number_format = '#,##,##0.00;[Red]-#,##,##0.00'
        for j, col in enumerate(cols, 1):
            longest = max([len(str(col))] + [len(str(r[j - 1])) for r in sec['rows'][:300]])
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, longest + 2), 60)
        ws.freeze_panes = ws.cell(r0 + 1, 1)
        ws.auto_filter.ref = f'A{r0}:{get_column_letter(len(cols))}{r0 + len(sec["rows"])}'
    if not wb.sheetnames:
        wb.create_sheet('Empty')
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def to_pdf(rep):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    bio = io.BytesIO()
    styles = getSampleStyleSheet()
    small = ParagraphStyle('s', parent=styles['Normal'], fontSize=7.4, leading=9)
    h1 = ParagraphStyle('h1', parent=styles['Title'], fontSize=16, textColor=colors.HexColor('#312E81'), alignment=0, spaceAfter=2)
    h2 = ParagraphStyle('h2', parent=styles['Heading2'], fontSize=11.5, textColor=colors.HexColor('#1F2937'), spaceBefore=10, spaceAfter=4)
    sub = ParagraphStyle('sub', parent=styles['Normal'], fontSize=8.5, textColor=colors.HexColor('#6B7280'))
    doc = SimpleDocTemplate(bio, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm,
                            bottomMargin=14 * mm, title=rep['title'], author='Finance Hub')
    W = landscape(A4)[0] - 24 * mm

    def esc(s):
        return str(s if s is not None else '').replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    def cell(v, col, money):
        if col in money and isinstance(v, (int, float)):
            return Paragraph(f'<para alignment="right">{inr(v, "")}</para>', small)
        return Paragraph(esc(v), small)

    story = [Paragraph(esc(rep['title']), h1), Paragraph(esc(rep.get('subtitle', '')), sub), Spacer(1, 4)]
    for sec in rep['sections']:
        story.append(Paragraph(esc(sec['title']), h2))
        if sec['kind'] == 'kv':
            data = [[Paragraph(f'<b>{esc(k)}</b>', small),
                     Paragraph(esc(inr(v, 'Rs ') if isinstance(v, (int, float)) and not isinstance(v, bool) else v), small)]
                    for k, v in sec['rows']]
            t = Table(data, colWidths=[75 * mm, 80 * mm], hAlign='LEFT')
            t.setStyle(TableStyle([('LINEBELOW', (0, 0), (-1, -1), 0.25, colors.HexColor('#E5E7EB')),
                                   ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
            story.append(t)
            continue
        cols, money = sec['columns'], set(sec.get('money', []))
        if not sec['rows']:
            story.append(Paragraph('Nothing to show.', sub))
            continue
        lens = [max([len(str(c))] + [min(len(str(r[i] if not isinstance(r[i], float) else inr(r[i], ''))), 60) for r in sec['rows'][:200]])
                for i, c in enumerate(cols)]
        total = sum(lens) or 1
        widths = [max(W * l / total, 14 * mm) for l in lens]
        scale = W / sum(widths)
        widths = [w * scale for w in widths]
        data = [[Paragraph(f'<b>{esc(c)}</b>', small) for c in cols]]
        data += [[cell(v, cols[i], money) for i, v in enumerate(r)] for r in sec['rows']]
        if sec.get('totals'):
            data.append([Paragraph('<b>Total</b>', small)] + [
                Paragraph(f'<para alignment="right"><b>{inr(sum(float(r[i] or 0) for r in sec["rows"] if isinstance(r[i], (int, float))), "")}</b></para>', small)
                if c in sec['totals'] else Paragraph('', small) for i, c in enumerate(cols) if i > 0])
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EEF2FF')),
            ('LINEBELOW', (0, 0), (-1, 0), 0.8, colors.HexColor('#6366F1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('LINEBELOW', (0, 1), (-1, -1), 0.2, colors.HexColor('#E5E7EB')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 2), ('BOTTOMPADDING', (0, 0), (-1, -1), 2)]))
        story.append(t)

    def footer(c, d):
        c.saveState()
        c.setFont('Helvetica', 7)
        c.setFillColor(colors.HexColor('#9CA3AF'))
        c.drawString(12 * mm, 7 * mm, f"Finance Hub · {rep['title']} · generated {datetime.now():%d %b %Y %H:%M} · confidential")
        c.drawRightString(landscape(A4)[0] - 12 * mm, 7 * mm, f'Page {d.page}')
        c.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    bio.seek(0)
    return bio


def send(rep, fmt, stem):
    stem = re.sub(r'[^A-Za-z0-9_-]+', '_', stem).strip('_')[:60] or 'report'
    if fmt == 'pdf':
        return send_file(to_pdf(rep), mimetype='application/pdf', as_attachment=True, download_name=f'{stem}.pdf')
    return send_file(to_xlsx(rep), as_attachment=True, download_name=f'{stem}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ── report builders ─────────────────────────────────────────────────────────
def scope_label():
    c = SC.current()
    return c['label'] or 'All books and accounts'


def view_report(con, start, end):
    """Money in / out for the current view: summary, by head, by month, accounts, transactions, linked transfers."""
    sql, par = SC.txn_sql('t')
    _, by_id = C.load_tree(con)
    rows = con.execute(f"SELECT t.*, c.kind AS kind FROM transactions t LEFT JOIN categories c ON c.id=t.category_id "
                       f"WHERE t.date>=? AND t.date<=? AND {sql} ORDER BY t.date, t.id", (start, end, *par)).fetchall()
    tot = defaultdict(float)
    heads = defaultdict(float)
    months = defaultdict(lambda: [0.0, 0.0])
    accts = defaultdict(lambda: [0.0, 0.0, 0])
    for r in rows:
        k = r['kind'] or 'expense'
        v = r['amount'] if r['type'] == 'debit' else -r['amount']
        if k == 'income':
            v = -v
        tot[k] += v
        n = by_id.get(r['category_id'])
        p = n['path'].split(' › ') if n else ['?']
        if k in ('income', 'expense'):
            heads[(k, p[1] if len(p) > 1 else p[0])] += v
        months[r['date'][:7]][0 if r['type'] == 'credit' else 1] += r['amount']
        a = accts[r['account']]
        a[0 if r['type'] == 'credit' else 1] += r['amount']
        a[2] += 1
    summary = [('View', scope_label()), ('Period', f'{start} to {end}'), ('Transactions', len(rows)),
               ('Income', round(tot['income'], 2)), ('Expenses', round(tot['expense'], 2)),
               ('Invested (net)', round(tot['investment'], 2)), ('Saved (income − expenses)', round(tot['income'] - tot['expense'], 2)),
               ('Transfers (not income / expense)', round(abs(tot['transfer']), 2))]
    secs = [{'title': 'Summary', 'kind': 'kv', 'rows': summary},
            {'title': 'By head', 'kind': 'table', 'columns': ['Type', 'Head', 'Amount'], 'money': ['Amount'],
             'rows': [('Income' if k == 'income' else 'Expense', h, round(v, 2)) for (k, h), v in sorted(heads.items(), key=lambda x: (x[0][0] != 'income', -x[1]))]},
            {'title': 'By month', 'kind': 'table', 'columns': ['Month', 'Money in', 'Money out', 'Net'], 'money': ['Money in', 'Money out', 'Net'],
             'rows': [(m, round(v[0], 2), round(v[1], 2), round(v[0] - v[1], 2)) for m, v in sorted(months.items())], 'totals': ['Money in', 'Money out', 'Net']},
            {'title': 'By account', 'kind': 'table', 'columns': ['Account', 'Money in', 'Money out', 'Transactions'], 'money': ['Money in', 'Money out'],
             'rows': [(a, round(v[0], 2), round(v[1], 2), v[2]) for a, v in sorted(accts.items())], 'totals': ['Money in', 'Money out']}]
    txr = []
    for r in rows:
        n = by_id.get(r['category_id'])
        txr.append((r['date'], r['account'], (r['display_name'] or r['payee'] or r['description'])[:70],
                    ' › '.join(n['path'].split(' › ')[1:]) if n else '', r['amount'] if r['type'] == 'debit' else '',
                    r['amount'] if r['type'] == 'credit' else '', r['entity'] or 'Personal', r['person'] or '', r['notes'] or ''))
    secs.append({'title': 'Transactions', 'kind': 'table', 'money': ['Out', 'In'], 'totals': ['Out', 'In'],
                 'columns': ['Date', 'Account', 'Description', 'Category', 'Out', 'In', 'Book', 'Person', 'Notes'], 'rows': txr})
    try:
        import ledger as L
        links = [l for l in L.link_rows(con, 'confirmed', 5000) if start <= (l['date'] or '') <= end]
        secs.append({'title': 'Linked transfers', 'kind': 'table', 'money': ['Amount'], 'totals': ['Amount'],
                     'columns': ['Date', 'Kind', 'From', 'To', 'Amount', 'Why'],
                     'rows': [(l['date'], l['label'], l['o_acct'] or l['counter'] or '', l['i_acct'] or l['counter'] or '', l['amount'] or 0, l['reason'] or '')
                              for l in links]})
    except Exception:  # noqa: BLE001
        pass
    return {'title': f'Statement — {scope_label()}', 'subtitle': f'{start} to {end}', 'sections': secs}


def trades_report(con, a):
    cond, par = ['1=1'], []
    bid = a.get('broker') or (str(SC.broker_id()) if SC.broker_id() else '')
    if bid:
        cond.append('t.broker_id=?'); par.append(int(bid))
    for k, col in (('start', 't.trade_date>=?'), ('end', 't.trade_date<=?'), ('side', 't.side=?')):
        if a.get(k):
            cond.append(col); par.append(a[k])
    if a.get('type'):
        cond.append("COALESCE(s.asset_type,'')=?"); par.append('' if a['type'] == 'none' else a['type'])
    if a.get('q'):
        cond.append('(s.symbol LIKE ? OR s.name LIKE ? OR s.isin LIKE ?)'); par += [f"%{a['q']}%"] * 3
    rows = con.execute(f"SELECT t.*, s.symbol, s.name, s.isin, s.asset_type, b.name AS broker FROM inv_trades t "
                       f"JOIN inv_securities s ON s.id=t.security_id JOIN inv_brokers b ON b.id=t.broker_id "
                       f"WHERE {' AND '.join(cond)} ORDER BY b.name, t.trade_date, t.id", par).fetchall()
    per = defaultdict(lambda: [0, 0.0, 0.0, 0.0])
    for r in rows:
        p = per[r['broker']]
        p[0] += 1
        p[1 if r['side'] == 'buy' else 2] += r['qty'] * r['price']
        p[3] += r['charges'] or 0
    import assets as AS
    groups = AS.broker_groups(con, int(bid) if bid else None)
    secs = [{'title': 'By broker', 'kind': 'table', 'columns': ['Broker', 'Trades', 'Bought', 'Sold', 'Charges', 'Net invested'],
             'money': ['Bought', 'Sold', 'Charges', 'Net invested'], 'totals': ['Bought', 'Sold', 'Charges', 'Net invested'],
             'rows': [(b, v[0], round(v[1], 2), round(v[2], 2), round(v[3], 2), round(v[1] - v[2] + v[3], 2)) for b, v in sorted(per.items())]},
            {'title': 'Trades', 'kind': 'table', 'money': ['Price', 'Value', 'Charges'], 'totals': ['Value', 'Charges'],
             'columns': ['Broker', 'Date', 'Security', 'ISIN', 'Type', 'Side', 'Qty', 'Price', 'Value', 'Charges'],
             'rows': [(r['broker'], r['trade_date'], r['name'] or r['symbol'], r['isin'] or '', r['asset_type'] or '', r['side'].upper(),
                       r['qty'], r['price'], round(r['qty'] * r['price'], 2), r['charges'] or 0) for r in rows]},
            {'title': 'Holdings by broker', 'kind': 'table', 'money': ['Avg cost', 'Price', 'Invested', 'Value', 'P&L'],
             'totals': ['Invested', 'Value', 'P&L'],
             'columns': ['Broker', 'Security', 'Type', 'Qty', 'Avg cost', 'Price', 'Invested', 'Value', 'P&L'],
             'rows': [(g['name'], r['name'], r['type_label'], r['qty'], round(r['avg'], 2), round(r['cmp'], 2), round(r['invested'], 2),
                       round(r['value'], 2), round(r['pnl'], 2)) for g in groups for r in g['rows']]}]
    who = next((g['name'] for g in groups), 'all brokers') if bid else 'all brokers'
    return {'title': f'Trades & holdings — {who}', 'subtitle': f"{a.get('start') or 'start'} to {a.get('end') or 'today'}", 'sections': secs}


def register(app, get_db):

    def report_view(fmt):
        con = get_db()
        fs, fe = fy_default()
        start, end = request.args.get('start') or fs, request.args.get('end') or fe
        if SC.current()['broker'] and not SC.current()['account'] and not SC.current()['book']:
            return send(trades_report(con, request.args), fmt, f"trades_{SC.current()['broker_name']}_{start}_{end}")
        return send(view_report(con, start, end), fmt, f"{scope_label()}_{start}_{end}")

    def report_trades(fmt):
        return send(trades_report(get_db(), request.args), fmt, f"trades_{date.today()}")

    def report_transfers(fmt):
        import ledger as L
        con = get_db()
        rows = L.link_rows(con, 'confirmed', 10000)
        flows = defaultdict(float)
        for r in rows:
            flows[(r['o_acct'] or r['counter'] or '?', r['i_acct'] or r['counter'] or '?', r['label'])] += r['amount'] or 0
        rep = {'title': f'Transfers between accounts — {scope_label()}', 'subtitle': 'confirmed links', 'sections': [
            {'title': 'Flows', 'kind': 'table', 'columns': ['From', 'To', 'Kind', 'Amount'], 'money': ['Amount'], 'totals': ['Amount'],
             'rows': [(a, b, k, round(v, 2)) for (a, b, k), v in sorted(flows.items(), key=lambda x: -x[1])]},
            {'title': 'Every link', 'kind': 'table', 'columns': ['Date', 'Kind', 'From', 'Out description', 'To', 'In description', 'Amount'],
             'money': ['Amount'], 'totals': ['Amount'],
             'rows': [(r['date'], r['label'], r['o_acct'] or r['counter'] or '', (r['o_desc'] or '')[:50], r['i_acct'] or r['counter'] or '',
                       (r['i_desc'] or '')[:50], r['amount'] or 0) for r in rows]}]}
        return send(rep, fmt, f'transfers_{date.today()}')

    def report_director(fmt):
        import ledger as L
        con = get_db()
        book = request.args.get('book') or next((b for b in SC.books(con) if b != 'Personal'), '')
        if not book:
            flash('Add your company as a book first.', 'info')
            return redirect(url_for('people_page'))
        start, end = request.args.get('start') or None, request.args.get('end') or None
        rows, opening = L.director_ledger(con, book, start, end)
        closing = opening + sum(r['effect'] for r in rows)
        rep = {'title': f'Director’s account — {book}', 'subtitle': f"{start or 'beginning'} to {end or 'today'}", 'sections': [
            {'title': 'Summary', 'kind': 'kv', 'rows': [('Opening balance', round(opening, 2)), ('Closing balance', round(closing, 2)),
                                                        ('Meaning', f'{book} owes you' if closing >= 0 else f'you owe {book}')]},
            {'title': 'Movements', 'kind': 'table', 'columns': ['Date', 'What', 'Detail', 'From', 'To', 'Amount', 'Effect', 'Balance'],
             'money': ['Amount', 'Effect', 'Balance'],
             'rows': [(r['date'], r['what'], r['detail'], r['from'], r['to'], r['amount'] or 0, round(r['effect'], 2), round(r['balance'], 2)) for r in rows]}]}
        return send(rep, fmt, f'director_account_{book}_{date.today()}')

    def report_reimb(fmt):
        import ledger as L
        con = get_db()
        ps = L.payers(con)
        payer = request.args.get('payer') or (ps[0] if ps else '')
        st = L.reimb_state(con, payer)
        rep = {'title': f'Reimbursements — {payer}', 'subtitle': f'as on {date.today()}', 'sections': [
            {'title': 'Summary', 'kind': 'kv', 'rows': [('Still to be paid back', round(st['outstanding'], 2)),
                                                        ('Not yet claimed', round(st['open_total'], 2)), ('Received', round(st['received_total'], 2))]},
            {'title': 'Claims', 'kind': 'table', 'columns': ['Claim', 'Date', 'Status', 'Expenses', 'Received', 'Balance'],
             'money': ['Expenses', 'Received', 'Balance'], 'totals': ['Expenses', 'Received', 'Balance'],
             'rows': [(c['title'], c['claim_date'], c['state'], round(c['total'], 2), round(c['received'], 2), round(c['balance'], 2)) for c in st['claims']]},
            {'title': 'Claimed expenses', 'kind': 'table', 'columns': ['Claim', 'Date', 'Account', 'Description', 'Amount'], 'money': ['Amount'],
             'totals': ['Amount'], 'rows': [(c['title'], i['date'], i['account'], (i['display_name'] or i['description'])[:60], i['amount'])
                                            for c in st['claims'] for i in c['items']] +
            [('(not claimed yet)', i['date'], i['account'], (i['display_name'] or i['description'])[:60], i['amount']) for i in st['open_items']]}]}
        return send(rep, fmt, f'reimbursements_{payer}_{date.today()}')

    def report_people_loans(fmt):
        import ledger as L
        con = get_db()
        loans = L.pl_rows(con, SC.current()['book'] or None)
        rep = {'title': f'Loans with people — {scope_label()}', 'subtitle': f'as on {date.today()}', 'sections': [
            {'title': 'Loans', 'kind': 'table', 'columns': ['Book', 'Person', 'Direction', 'Started', 'Due', 'Amount', 'Repaid', 'Outstanding', 'Status'],
             'money': ['Amount', 'Repaid', 'Outstanding'], 'totals': ['Amount', 'Repaid', 'Outstanding'],
             'rows': [(l['book'], l['person'], 'You lent' if l['direction'] == 'given' else 'You borrowed', l['start_date'], l['due_date'] or '',
                       round(l['principal'], 2), round(l['repaid'], 2), round(l['outstanding'], 2), l['status']) for l in loans]},
            {'title': 'Movements', 'kind': 'table',
             'columns': ['Person', 'Date', 'What', 'Amount', 'Mode', 'Our account', 'Their account', 'IFSC / UPI', 'Reference'],
             'money': ['Amount'],
             'rows': [(l['person'], m['date'], m['kind'], m['amount'], m['mode'] or '', m['our_account'] or '', m['their_account'] or '',
                       m['their_ifsc'] or m['their_upi'] or '', m['reference'] or '') for l in loans for m in l['moves']]}]}
        return send(rep, fmt, f'loans_with_people_{date.today()}')

    for rule, ep, fn in [
        ('/report/view/<fmt>', 'report_view', report_view), ('/report/trades/<fmt>', 'report_trades', report_trades),
        ('/report/transfers/<fmt>', 'report_transfers', report_transfers), ('/report/director/<fmt>', 'report_director', report_director),
        ('/report/reimbursements/<fmt>', 'report_reimb', report_reimb), ('/report/people-loans/<fmt>', 'report_people_loans', report_people_loans),
    ]:
        app.add_url_rule(rule, ep, fn, methods=['GET'])
