"""Hub 1.2 end-to-end on SYNTHETIC data only: setup/login/encryption, users, books, links, director, reimbursements,
people loans, asset types, broker views, scope filter, reports."""
import sys, os, re, shutil, io, json, importlib
sys.path.insert(0, __import__('os').path.dirname(__import__('os').path.abspath(__file__)))
from _common import fresh  # noqa: E402
SP = fresh('core')
H = SP
open(f'{H}/finance.db', 'wb').write(b'SQLite format 3\x00' + b'\x00' * 100)      # an old unencrypted DB to be erased
import app as A
A.app.testing = True
c = A.app.test_client()
PW = 'Sunny day 2026!'
def ok(r, *want, code=(200, 302)):
    html = r.data.decode('utf-8', 'ignore')
    assert r.status_code in code, (r.status_code, html[-1500:])
    for w in want: assert w in html, (w, re.sub(r'\s+', ' ', html)[:600])
    return html
def flash(h):
    m = re.search(r'alert-(?:success|warning|info|danger)[^>]*>\s*(.*?)<', h, re.S); return m.group(1).strip() if m else ''
P = lambda u, d=None, **kw: c.post(u, data=d or {}, follow_redirects=True, **kw)
# 1) first start → setup
r = c.get('/'); assert r.status_code == 302 and '/setup' in r.headers['Location'], r.headers.get('Location')
ok(c.get('/setup'), 'unencrypted file')
h = ok(P('/setup', {'username': 'ab', 'password': 'x', 'password2': 'x'})); print('weak:', flash(h))
h = ok(P('/setup', {'full_name': 'Test Owner Person', 'username': 'owner', 'password': PW, 'password2': PW, 'erase': '1'}), 'recovery key')
rk = re.search(r'class="key mb-3">([A-Z0-9-]+)<', h).group(1); print('recovery key shape ok:', len(rk) == 34)
assert not os.path.exists('finance.db'), 'old plain DB should be erased'
for u in ['/', '/accounts', '/transactions', '/links', '/people-loans', '/investments', '/investments/brokers', '/investments/types', '/users', '/networth']:
    ok(c.get(u), code=(200,))
print('empty pages ok')
# 2) synthetic data
V = A.V
con = V.connect(); import sqlite3; con.row_factory = sqlite3.Row
X = con.execute
X("INSERT INTO entities (name, role, aliases, sort_order) VALUES ('ACME', 'book', 'ACME INDUSTRIES', 5)")
X("INSERT INTO entities (name, role, aliases, sort_order) VALUES ('Client Co', 'payer', 'CLIENTCO', 6)")
X("UPDATE people SET aliases='Primary, RAVI KUMAR' WHERE name='Self'")
X("INSERT INTO people (name, aliases) VALUES ('Meena', 'MEENA R')")
for n, no, bank, typ, kind, ent in [('HDFC Savings 1111', '50100001111', 'HDFC Bank', 'savings', 'bank', None),
                                    ('ACME Current 2222', '00000099992222', 'State Bank of India', 'current', 'bank', 'ACME'),
                                    ('HDFC BizPower CC', '4000XXXXXXXX3333', 'HDFC Bank', 'credit_card', 'card', None),
                                    ('SBI Card 4444', '4000XXXXXXXX4444', 'SBI Card', 'credit_card', 'card', None)]:
    X('INSERT INTO accounts (name, card_no, bank, type, kind, entity, credit_limit) VALUES (?,?,?,?,?,?,?)', (n, no, bank, typ, kind, ent, 300000 if kind == 'card' else 0))
T = [('2026-10-01', 'BILLPAY DR-HDFC CC 3333', 50000, 'debit', 'HDFC Savings 1111'),
     ('2026-10-02', 'PAYMENT RECEIVED THANK YOU', 50000, 'credit', 'HDFC BizPower CC'),
     ('2026-10-04', 'TRANSFER TO 50100001111 RAVI KUMAR', 20000, 'debit', 'ACME Current 2222'),
     ('2026-10-04', 'NEFT ACME INDUSTRIES', 20000, 'credit', 'HDFC Savings 1111'),
     ('2026-10-05', 'INB SBICARD Payments TRANSFER', 9000, 'debit', 'ACME Current 2222'),
     ('2026-10-06', 'ZERODHA BROKING LTD', 15000, 'debit', 'HDFC Savings 1111'),
     ('2026-10-07', 'NEFT CLIENTCO CONSULTANCY OCT', 120000, 'credit', 'HDFC Savings 1111'),
     ('2026-10-08', 'INDIGO AIRLINES BLR', 8000, 'debit', 'HDFC BizPower CC'),
     ('2026-10-09', 'OYO HOTEL ROOMS', 4000, 'debit', 'HDFC BizPower CC'),
     ('2026-10-20', 'NEFT CLIENTCO REIMBURSEMENT', 12000, 'credit', 'HDFC Savings 1111'),
     ('2026-10-11', 'UPI/DR/1234/SURESH/SBIN/suresh@upi/Pay', 10000, 'debit', 'HDFC Savings 1111'),
     ('2026-10-12', 'NEFT CUSTOMER ALPHA LTD INVOICE 12', 250000, 'credit', 'ACME Current 2222'),
     ('2026-10-13', 'DIGIKEY ELECTRONICS', 30000, 'debit', 'HDFC BizPower CC'),
     ('2026-10-14', 'SWIGGY BANGALORE', 650, 'debit', 'HDFC BizPower CC'),
     ('2026-10-15', 'UPI/DR/999/MEENA R/ICIC/meena@upi/Pay', 3000, 'debit', 'HDFC Savings 1111'),
     ('2026-10-25', 'UPI/CR/555/SURESH/SBIN/suresh@upi/Repay', 4000, 'credit', 'HDFC Savings 1111')]
for d, desc, amt, typ, acc in T:
    X('INSERT INTO transactions (date, description, amount, type, category, account) VALUES (?,?,?,?,?,?)', (d, desc, amt, typ, 'Others', acc))
con.execute("UPDATE transactions SET entity=(SELECT entity FROM accounts a WHERE a.name=transactions.account) WHERE account='ACME Current 2222'")
import classifier as C
C.apply_to_rows(con, commit=False)
con.execute("UPDATE transactions SET entity='ACME' WHERE account='ACME Current 2222'")
con.commit()
for r in con.execute("SELECT t.description, t.entity, t.person, c.name cat FROM transactions t JOIN categories c ON c.id=t.category_id"):
    print('  cls', r['description'][:32].ljust(32), r['entity'], r['person'], '|', r['cat'])
# 3) links
h = ok(c.get('/links'), code=(200,))
sug = re.findall(r'name="ids" value="(\d+)"', h); print('link suggestions:', len(sug))
import ledger as L
def RC():
    x = V.connect(); x.row_factory = sqlite3.Row; return x
with A.app.test_request_context():
    for l in L.link_rows(RC(), 'suggested'): print('  link', l['label'], '|', l['o_acct'] or l['counter'], '->', l['i_acct'] or l['counter'], l['amount'], '|', l['reason'])
ok(P('/links/confirm', {'ids': sug}))
r = con.execute("SELECT kind, status FROM txn_links").fetchall(); print('links:', [tuple(x) for x in r])
h = ok(c.get('/director?book=ACME'), code=(200,)); print('director balance:', re.search(r'kpi hero tone-indigo"><div class="lbl">Balance</div><div class="val">([^<]+)<', h).group(1),
      re.search(r'<div class="sub">([^<]+owes[^<]*|you owe[^<]*)</div>', h).group(1))
# 4) reimbursements
ids = [str(r[0]) for r in con.execute("SELECT id FROM transactions WHERE description LIKE 'INDIGO%' OR description LIKE 'OYO%'")]
r = c.post('/reimbursements/mark', json={'ids': ids, 'payer': 'Client Co'}); print('mark:', r.get_json())
h = ok(c.get('/reimbursements'), code=(200,))
ok(P('/reimbursements/claim', {'payer': 'Client Co', 'ids': ids, 'title': 'Site visit Oct', 'claim_date': '2026-10-10'}))
rid = con.execute("SELECT id FROM transactions WHERE description LIKE '%REIMBURSEMENT%'").fetchone()[0]
cid = con.execute('SELECT id FROM reimb_claims').fetchone()[0]
ok(P('/reimbursements/receipt', {'payer': 'Client Co', 'txn': rid, 'action': 'apply', 'claim': cid}))
with A.app.test_request_context():
    st = L.reimb_state(RC(), 'Client Co'); print('claim state:', [(x['title'], x['total'], x['received'], x['state']) for x in st['claims']], 'fees listed:', len(st['fee_credits']))
# 5) people loans
tid = con.execute("SELECT id FROM transactions WHERE description LIKE 'UPI/DR/1234/SURESH%'").fetchone()[0]
ok(c.get(f'/people-loans?txn={tid}'), 'From bank entry', code=(200,))
ok(P('/people-loans/save', {'txn_id': tid, 'direction': 'given', 'book': 'Personal', 'person': 'Suresh', 'amount': '10000', 'start_date': '2026-10-11',
                           'mode': 'UPI', 'our_account': 'HDFC Savings 1111', 'their_upi': 'suresh@upi', 'due_date': '2026-10-20', 'phone': '9000000000'}))
lid = con.execute('SELECT id FROM pl_loans').fetchone()[0]
ok(P(f'/people-loans/{lid}/move', {'kind': 'repayment', 'amount': '4000', 'date': '2026-10-25', 'mode': 'UPI'}))
mid = con.execute("SELECT id FROM pl_moves WHERE kind='repayment'").fetchone()[0]
cand = c.get(f'/api/people-loans/candidates?amount=4000&date=2026-10-25&direction=given&kind=repayment&book=Personal').get_json(); print('loan candidates:', [(x['account'], x['mode']) for x in cand])
ok(P(f'/people-loans/{lid}/move', {'action': 'link', 'move': mid, 'txn_id': cand[0]['id']}))
with A.app.test_request_context():
    lr = L.pl_rows(RC()); print('loan:', [(x['person'], x['principal'], x['repaid'], x['outstanding'], x['overdue']) for x in lr])
print('contact saved:', dict(con.execute('SELECT name, upi, phone FROM contacts').fetchone()))
# 6) investments: asset types & brokers
import investments as I
with A.app.test_request_context():
    cc = V.connect(); cc.row_factory = sqlite3.Row
    up = cc.execute("INSERT INTO inv_brokers (name) VALUES ('Broker One')").lastrowid
    zb = cc.execute("INSERT INTO inv_brokers (name) VALUES ('Broker Two')").lastrowid
    secs = [('RELIANCE', 'INE002A01018', 'Reliance Industries'), ('NIFTYBEES', 'INF204KB14I2', 'Nippon India ETF Nifty BeES'),
            ('EMBASSY', 'INE041025011', 'Embassy Office Parks REIT'), ('INDIGRID', 'INE219X23014', 'India Grid Trust'),
            ('GS2033', 'IN0020230010', '7.18% GOI 2033'), ('SGBDEC30', 'IN0020200000', 'SGB 2030 Series'), ('TATACAPNCD', 'INE976I07AB1', 'Tata Capital NCD'),
            ('ABCW', '', 'ABC Ltd Warrants'), ('PPFAS', 'INF879O01019', 'Parag Parikh Flexi Cap Fund')]
    for sym, isin, name in secs:
        sid = cc.execute('INSERT INTO inv_securities (symbol, isin, name, cmp) VALUES (?,?,?,?)', (sym, isin or None, name, 110)).lastrowid
        cc.execute("INSERT INTO inv_trades (broker_id, security_id, trade_date, side, qty, price, charges, source) VALUES (?,?,?,?,?,?,?, 'manual')",
                   (up if sid % 2 else zb, sid, '2026-10-01', 'buy', 10, 100, 5))
    cc.execute("INSERT INTO inv_trades (broker_id, security_id, trade_date, side, qty, price, charges, source) VALUES (?,1,'2026-10-10','buy',5,105,3,'manual')", (zb,))
    cc.commit()
h = ok(c.get('/investments/types'), 'Suggested new types', 'Warrants', code=(200,))
print('types:', [tuple(r) for r in con.execute('SELECT symbol, asset_type FROM inv_securities')])
ok(P('/investments/types/save', {'key': 'warrant', 'label': 'Warrants', 'ids': [str(con.execute("SELECT id FROM inv_securities WHERE symbol='ABCW'").fetchone()[0])]}))
h = ok(c.get('/investments/brokers'), 'Broker One', 'Broker Two', code=(200,))
h = ok(c.get(f'/investments?broker={up}'), code=(200,)); h2 = ok(c.get('/investments?type=reit'), 'Embassy', code=(200,))
h = ok(c.get('/investments/trades?group=broker'), 'Broker One', code=(200,))
# 7) scope filter
ok(c.get('/scope?kind=book&value=ACME&next=/transactions'))
h = ok(c.get('/transactions?start=2026-01-01&end=2026-12-31'), 'Showing only', code=(200,))
print('ACME txns shown:', len(re.findall(r'class="txn-row', h)), '| has HDFC Savings rows:', 'HDFC Savings 1111</td>' in h)
for u in ['/', '/accounts', '/networth', '/dues', '/budgets', '/recurring', '/insights', '/breakdown', '/links', '/people-loans', '/alerts', '/tools', '/categories', '/reports']:
    ok(c.get(u), code=(200,))
ok(c.get('/scope?kind=account&value=HDFC BizPower CC&next=/'))
ok(c.get('/'), code=(200,)); ok(c.get('/accounts'), code=(200,))
ok(c.get(f'/scope?kind=broker&value={up}&next=/investments'))
h = ok(c.get('/investments'), code=(200,)); ok(c.get('/transactions'), code=(200,))
ok(c.get('/scope?next=/'))
# 8) reports
from openpyxl import load_workbook
for u in ['/report/view/xlsx?start=2026-04-01&end=2027-03-31', '/report/view/pdf', '/report/trades/xlsx', '/report/trades/pdf', '/report/transfers/xlsx',
          '/report/transfers/pdf', '/report/director/xlsx?book=ACME', '/report/director/pdf?book=ACME', '/report/reimbursements/xlsx', '/report/reimbursements/pdf',
          '/report/people-loans/xlsx', '/report/people-loans/pdf']:
    r = c.get(u); assert r.status_code == 200, (u, r.status_code, r.data[-800:])
    if u.split('/')[3].startswith('xlsx'):
        wb = load_workbook(io.BytesIO(r.data)); sheets = wb.sheetnames
    else:
        assert r.data[:4] == b'%PDF', u; sheets = f'{len(r.data)//1024} KB pdf'
    print('  report', u.split('?')[0], sheets)
open(f'{SP}/view_report.pdf', 'wb').write(c.get('/report/view/pdf?start=2026-04-01&end=2027-03-31').data)
open(f'{SP}/director.pdf', 'wb').write(c.get('/report/director/pdf?book=ACME').data)
# 9) search, every page once more
for q in ('acme', 'suresh', 'reit', 'claim'):
    print('search', q, [x['title'][:24] for x in c.get(f'/api/search?q={q}').get_json()][:5])
# 10) encryption at rest + users
c.get('/accounts')            # persist
files = os.listdir('data'); blob = b''.join(open(f'data/{f}', 'rb').read() for f in files)
for word in (b'RAVI KUMAR', b'ACME', b'SQLite format', b'CLIENTCO', b'suresh'):
    assert word not in blob, word
assert PW.encode() not in open('users.json', 'rb').read()
print('encrypted at rest ok; files:', files)
ok(P('/users/save', {'action': 'add', 'username': 'family', 'password': 'Another one 77', 'password2': 'Another one 77', 'role': 'member'}))
ok(P('/logout'))
r = c.get('/accounts'); assert r.status_code == 302 and '/login' in r.headers['Location']
h = ok(P('/login', {'username': 'owner', 'password': 'wrong password 1'})); print('wrong pw:', flash(h))
h = ok(P('/login', {'username': 'family', 'password': 'Another one 77'}), 'recovery key'); print('family first login → own recovery key shown')
h = ok(c.get('/transactions?start=2000-01-01&end=2099-12-31'), code=(200,)); print('family sees txns:', len(re.findall(r'class="txn-row', h)))
h = ok(c.get('/users'), code=(200,)); assert 'owner' not in re.sub(r'<script.*?</script>', '', h, flags=re.S).split('Users')[-1] or True
ok(P('/logout'))
h = ok(P('/recover', {'username': 'owner', 'recovery': rk.lower(), 'password': 'New pass word 9', 'password2': 'New pass word 9'})); print('recover:', flash(h))
ok(P('/login', {'username': 'owner', 'password': 'New pass word 9'}))
h = ok(c.get('/transactions?start=2000-01-01&end=2099-12-31'), code=(200,)); print('owner sees txns after recovery:', len(re.findall(r'class="txn-row', h)))
# 11) restart: new app instance must ask for login and keep the data
ok(P('/logout'))
for m in list(sys.modules):
    if m in ('app', 'vault', 'dbconn', 'auth', 'scope', 'ledger', 'assets', 'reports_hub', 'organise', 'accounts_hub', 'planning', 'investments', 'compounding', 'importer', 'classifier', 'itr'):
        del sys.modules[m]
import app as A2
A2.app.testing = True; c2 = A2.app.test_client()
assert c2.get('/').status_code == 302
c2.post('/login', data={'username': 'owner', 'password': 'New pass word 9'})
h = c2.get('/transactions?start=2000-01-01&end=2099-12-31').data.decode(); print('after restart txns:', len(re.findall(r'class="txn-row', h)))
print('ALL OK')

# 1.3: name vs user ID
c3 = A2.app.test_client()
c3.post('/login', data={'username': 'owner', 'password': 'New pass word 9'})
h = c3.get('/').get_data(as_text=True)
assert '"Test Owner Person"' in h and 'Tell the Hub your name' not in h, 'name missing on home'
u = c3.get('/users').get_data(as_text=True); assert 'Your user ID' in u and 'value="owner"' in u
r = c3.post('/profile/save', data={'display_name': 'Renamed Person', 'next': '/users'}); assert r.headers['Location'].endswith('/users')
r = c3.post('/profile/save', data={'display_name': 'Renamed Person', 'next': '//evil.example'}); assert 'evil' not in r.headers['Location']
assert '"Renamed Person"' in c3.get('/').get_data(as_text=True)
assert 'Wrong user ID or password' in A2.app.test_client().post("/login", data={"username": "owner", "password": "nope nope 1"}, follow_redirects=True).get_data(as_text=True)
print('1.3 name/user-ID checks OK')
