"""Parser for Central Bank of India (CBI) Savings Account XLSX statements."""
import openpyxl
from datetime import datetime
from pathlib import Path
from .categorizer import categorize

ACCOUNT = 'CBI Savings Account'


def normalize_date(raw) -> str:
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d %b %Y', '%d/%m/%y'):
        try:
            if isinstance(raw, str):
                return datetime.strptime(raw.strip(), fmt).strftime('%Y-%m-%d')
            elif hasattr(raw, 'strftime'):
                return raw.strftime('%Y-%m-%d')
        except (ValueError, AttributeError):
            pass
    return str(raw)


def parse_cbi_xlsx(filepath: str) -> list[dict]:
    transactions = []
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active

        header_row = None
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            row_str = ' '.join(str(c) for c in row if c).lower()
            if 'date' in row_str and ('debit' in row_str or 'withdrawal' in row_str or 'amount' in row_str):
                header_row = i
                headers = [str(c).lower().strip() if c else '' for c in row]
                break

        if header_row is None:
            return transactions

        date_idx = next((i for i, h in enumerate(headers) if 'date' in h), 0)
        desc_idx = next((i for i, h in enumerate(headers) if 'narration' in h or 'description' in h or 'particular' in h or 'remark' in h), 2)
        debit_idx = next((i for i, h in enumerate(headers) if 'debit' in h or 'withdrawal' in h or 'dr' in h), None)
        credit_idx = next((i for i, h in enumerate(headers) if 'credit' in h or 'deposit' in h or 'cr' in h), None)
        amount_idx = next((i for i, h in enumerate(headers) if 'amount' in h), None)

        for row in ws.iter_rows(min_row=header_row + 2, values_only=True):
            if not row or all(c is None for c in row):
                continue
            try:
                date_str = normalize_date(row[date_idx])
                desc = str(row[desc_idx]).strip() if row[desc_idx] else ''
                if not desc or desc == 'None':
                    continue

                amount = 0.0
                txn_type = 'debit'

                if debit_idx is not None and credit_idx is not None:
                    dr = row[debit_idx]
                    cr = row[credit_idx]
                    if dr and dr != 0:
                        try:
                            amount = float(str(dr).replace(',', ''))
                            txn_type = 'debit'
                        except ValueError:
                            pass
                    elif cr and cr != 0:
                        try:
                            amount = float(str(cr).replace(',', ''))
                            txn_type = 'credit'
                        except ValueError:
                            pass
                elif amount_idx is not None:
                    try:
                        amount = abs(float(str(row[amount_idx]).replace(',', '')))
                    except (ValueError, TypeError):
                        pass

                if amount == 0:
                    continue

                transactions.append({
                    'date': date_str,
                    'description': desc,
                    'amount': amount,
                    'type': txn_type,
                    'category': categorize(desc),
                    'account': ACCOUNT,
                    'source_file': Path(filepath).name,
                })
            except Exception:
                continue
    except Exception:
        pass
    return transactions


def load_all_cbi(folder: str) -> list[dict]:
    all_txns = []
    p = Path(folder)
    for f in sorted(p.glob('AccountStatement__*.xlsx')):
        all_txns.extend(parse_cbi_xlsx(str(f)))
    return all_txns
