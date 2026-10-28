"""Parser for HDFC Bank Credit Card XLS statements (BizPower & TataNeu).

Both card types share the same XLS format but have slightly different column positions:
  BizPower (card 3081): header at row 25, date col 11, desc col 16, amt col 27, type col 32
  TataNeu  (card 1952): header at row 28, date col 16, desc col 20, amt col 47, type col 52

Conversion: XLS → XLSX via LibreOffice subprocess (no password needed for HDFC cards).
"""
import re
import subprocess
import tempfile
import os
from datetime import datetime
from pathlib import Path
import openpyxl
from .categorizer import categorize

HDFC_BIZPOWER_ACCOUNT = 'HDFC BizPower CC'
HDFC_TATANEU_ACCOUNT = 'HDFC TataNeu CC'


def convert_xls_to_xlsx(xls_path: str) -> str | None:
    """Convert an XLS file to XLSX using LibreOffice. Returns path to XLSX or None."""
    out_dir = tempfile.mkdtemp()
    try:
        result = subprocess.run(
            ['libreoffice', '--headless', '--convert-to', 'xlsx', '--outdir', out_dir, xls_path],
            capture_output=True, text=True, timeout=60
        )
        stem = Path(xls_path).stem
        out_file = os.path.join(out_dir, stem + '.xlsx')
        if os.path.exists(out_file):
            return out_file
    except Exception:
        pass
    return None


def parse_date(raw) -> str:
    """Parse HDFC date formats: DD/MM/YYYY, DD/MM/YYYY HH:MM:SS, DD/MM/YY."""
    if raw is None:
        return ''
    s = str(raw).strip()
    # Strip time component if present
    s = s.split(' ')[0].split('\n')[0]
    for fmt in ('%d/%m/%Y', '%d/%m/%y', '%d-%m-%Y', '%d-%m-%y'):
        try:
            return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except ValueError:
            pass
    return s


def clean_amount(raw) -> float:
    """Clean amount: strip ₹, commas, whitespace."""
    if raw is None:
        return 0.0
    s = str(raw).replace('₹', '').replace(',', '').strip()
    try:
        return abs(float(s))
    except ValueError:
        return 0.0


def detect_hdfc_layout(ws) -> dict | None:
    """Scan worksheet to find header row and return column mapping.

    Returns dict with keys: header_row (1-based), date_col, desc_col, amt_col, type_col
    or None if not detected.
    """
    for i, row in enumerate(ws.iter_rows(min_row=20, max_row=40, values_only=True), start=20):
        cells = list(row)
        row_str = ' '.join(str(c) for c in cells if c)
        if 'DATE' not in row_str:
            continue
        # Find DATE column
        date_col = None
        for j, c in enumerate(cells):
            if c and str(c).strip().upper() == 'DATE':
                date_col = j
                break
        if date_col is None:
            continue
        # BizPower: DATE at col 11, TataNeu: DATE at col 16
        if date_col == 11:
            return {
                'header_row': i, 'data_row': i + 1,
                'date_col': 11, 'desc_col': 16, 'amt_col': 27, 'type_col': 32,
                'account': HDFC_BIZPOWER_ACCOUNT,
            }
        elif date_col == 16:
            return {
                'header_row': i, 'data_row': i + 1,
                'date_col': 16, 'desc_col': 20, 'amt_col': 47, 'type_col': 52,
                'account': HDFC_TATANEU_ACCOUNT,
            }
    return None


def parse_hdfc_xlsx(xlsx_path: str, source_filename: str) -> list[dict]:
    """Parse a converted HDFC CC XLSX and return transaction list."""
    transactions = []
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws = wb.active
        layout = detect_hdfc_layout(ws)
        if layout is None:
            return transactions

        date_col = layout['date_col']
        desc_col = layout['desc_col']
        amt_col = layout['amt_col']
        type_col = layout['type_col']
        account = layout['account']
        data_start = layout['data_row']

        for row in ws.iter_rows(min_row=data_start, values_only=True):
            cells = list(row)
            if not cells or all(c is None for c in cells):
                continue

            date_val = cells[date_col] if date_col < len(cells) else None
            if not date_val:
                continue
            date_str = parse_date(date_val)
            if not date_str or not re.match(r'\d{4}-\d{2}-\d{2}', date_str):
                continue

            desc_raw = cells[desc_col] if desc_col < len(cells) else ''
            desc = str(desc_raw).replace('\n', ' ').strip() if desc_raw else ''
            if not desc or desc.lower() in ('description', 'nan', 'none'):
                continue

            amount = clean_amount(cells[amt_col] if amt_col < len(cells) else None)
            if amount == 0:
                continue

            # Type: 'Cr' = credit/payment, blank = debit
            type_raw = str(cells[type_col]).strip() if type_col < len(cells) and cells[type_col] else ''
            txn_type = 'credit' if 'cr' in type_raw.lower() else 'debit'

            transactions.append({
                'date': date_str,
                'description': desc,
                'amount': amount,
                'type': txn_type,
                'category': categorize(desc),
                'account': account,
                'source_file': source_filename,
            })
    except Exception:
        pass
    return transactions


def parse_hdfc_xls(xls_path: str) -> list[dict]:
    """Convert XLS → XLSX then parse."""
    xlsx_path = convert_xls_to_xlsx(xls_path)
    if xlsx_path is None:
        return []
    try:
        return parse_hdfc_xlsx(xlsx_path, Path(xls_path).name)
    finally:
        try:
            os.unlink(xlsx_path)
            os.rmdir(os.path.dirname(xlsx_path))
        except Exception:
            pass


def load_all_hdfc_cc(bizpower_folder: str, tataneu_folder: str) -> list[dict]:
    """Load all HDFC CC XLS files from BizPower and TataNeu folders."""
    all_txns = []
    for folder in (bizpower_folder, tataneu_folder):
        p = Path(folder)
        if not p.exists():
            continue
        for f in sorted(p.glob('*.xls')):
            all_txns.extend(parse_hdfc_xls(str(f)))
        for f in sorted(p.glob('*.XLS')):
            all_txns.extend(parse_hdfc_xls(str(f)))

    # Deduplicate
    seen = set()
    unique = []
    for t in all_txns:
        key = (t['date'], t['description'][:50], t['amount'], t['account'])
        if key not in seen:
            seen.add(key)
            unique.append(t)
    return unique
